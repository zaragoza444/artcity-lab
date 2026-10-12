"""Consent, suppression, imports, and segment eligibility."""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import datetime, timedelta, timezone
from typing import Any

from sqlalchemy.orm import Session

from app.communications.audit import audit
from app.communications.foundation import utcnow
from app.communications.phones import normalize_phone
from app.core.config import get_settings
from app.models_comm import (
    CommConsent,
    CommContact,
    CommImportJob,
    CommMessage,
    CommRecipient,
    CommSuppression,
)

_SAMPLE_CAP = 200


def _purpose_matches(stored: str, requested: str) -> bool:
    if requested == "marketing":
        return stored in {"marketing", "event_updates", "promotion", "events"}
    return stored == requested


def active_suppression(db: Session, contact_id: str, now: datetime | None = None) -> CommSuppression | None:
    moment = now or utcnow()
    rows = db.query(CommSuppression).filter(CommSuppression.contact_id == contact_id).all()
    for row in rows:
        if row.expires_at is not None and row.expires_at <= moment:
            continue
        return row
    return None


def has_verified_consent(
    db: Session,
    contact_id: str,
    *,
    business_unit_id: str | None,
    purpose: str,
) -> bool:
    rows = (
        db.query(CommConsent)
        .filter(
            CommConsent.contact_id == contact_id,
            CommConsent.channel == "whatsapp",
            CommConsent.status == "granted",
            CommConsent.withdrawn_at.is_(None),
        )
        .all()
    )
    for row in rows:
        if not _purpose_matches(row.purpose, purpose):
            continue
        if business_unit_id and row.business_unit_id not in (None, business_unit_id):
            continue
        if row.obtained_at and row.source and row.proof_uri:
            return True
        if get_settings().comm_allow_unverified_consent and row.obtained_at:
            return True
    return False


def in_marketing_cooldown(db: Session, contact_id: str, now: datetime | None = None) -> bool:
    hours = get_settings().comm_marketing_cooldown_hours
    if hours <= 0:
        return False
    moment = now or utcnow()
    since = moment - timedelta(hours=hours)
    row = (
        db.query(CommMessage)
        .filter(
            CommMessage.contact_id == contact_id,
            CommMessage.direction == "outbound",
            CommMessage.accepted_at.is_not(None),
            CommMessage.accepted_at >= since,
        )
        .first()
    )
    return row is not None


def suppress_contact(
    db: Session,
    contact: CommContact,
    *,
    reason: str,
    actor_id: str | None,
    source: str,
) -> None:
    if active_suppression(db, contact.id) is None:
        db.add(
            CommSuppression(
                contact_id=contact.id,
                scope="global",
                reason=reason[:255],
                created_by=actor_id,
            )
        )
    now = utcnow()
    consents = (
        db.query(CommConsent)
        .filter(CommConsent.contact_id == contact.id, CommConsent.withdrawn_at.is_(None), CommConsent.status == "granted")
        .all()
    )
    for consent in consents:
        consent.withdrawn_at = now
        consent.status = "withdrawn"
    waiting = (
        db.query(CommRecipient)
        .filter(
            CommRecipient.contact_id == contact.id,
            CommRecipient.status.in_(["eligible", "queued", "retryable", "selected"]),
        )
        .all()
    )
    for recipient in waiting:
        recipient.status = "cancelled"
        recipient.exclusion_reason = "opted_out"
        recipient.updated_at = now
    audit(db, actor_id=actor_id, action="contact.opt_out", object_type="contact", object_id=contact.id, metadata={"source": source})
    db.flush()


def _parse_time(value: str | None) -> datetime | None:
    if not value:
        return None
    text = value.strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def classify_consent(row: dict[str, str]) -> str:
    status = (row.get("consent_status") or "").strip().lower()
    has_proof = bool(row.get("consent_timestamp") and row.get("consent_source") and row.get("consent_proof_ref"))
    if status in {"withdrawn", "opted_out", "denied", "stop"}:
        return "withdrawn"
    if status in {"granted", "opted_in", "yes"} and has_proof and _parse_time(row.get("consent_timestamp")):
        return "granted"
    if status in {"granted", "opted_in", "yes"}:
        return "unconfirmed"
    return "pending"


def _iter_rows(filename: str, data: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            return []
        headers = [str(cell or "").strip().lower().replace(" ", "_") for cell in header_row]
        rows = []
        for raw in iterator:
            if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
                continue
            item = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = raw[idx] if idx < len(raw) else ""
                item[header] = "" if value is None else str(value).strip()
            rows.append(item)
        book.close()
        return rows
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for raw in reader:
        rows.append({(key or "").strip().lower().replace(" ", "_"): (value or "").strip() for key, value in raw.items()})
    return rows


def import_contacts(
    db: Session,
    *,
    organization_id: str,
    business_unit_id: str | None,
    actor_id: str | None,
    filename: str,
    data: bytes,
) -> CommImportJob:
    checksum = hashlib.sha256(data).hexdigest()
    job = CommImportJob(
        organization_id=organization_id,
        business_unit_id=business_unit_id,
        imported_by=actor_id,
        checksum=checksum,
        status="running",
        rejections_json={"counts": {}, "samples": []},
    )
    db.add(job)
    db.flush()
    rows = _iter_rows(filename, data)
    counts: dict[str, int] = {}
    samples: list[dict[str, str]] = []
    seen: set[str] = set()
    inserted = updated = rejected = valid_opt_in = needs_consent = 0

    def reject(reason: str, phone_hint: str) -> None:
        nonlocal rejected
        rejected += 1
        counts[reason] = counts.get(reason, 0) + 1
        if len(samples) < _SAMPLE_CAP:
            samples.append({"reason": reason, "phone_hint": phone_hint[:8]})

    for row in rows:
        e164, iso, error = normalize_phone(row.get("phone_number") or row.get("phone") or "")
        if error or not e164:
            reject(error or "invalid_phone", row.get("phone_number") or "")
            continue
        if e164 in seen:
            reject("duplicate_in_file", e164)
            continue
        seen.add(e164)
        contact = (
            db.query(CommContact)
            .filter(CommContact.organization_id == organization_id, CommContact.phone_e164 == e164, CommContact.deleted_at.is_(None))
            .first()
        )
        created = contact is None
        if contact is None:
            contact = CommContact(
                organization_id=organization_id,
                phone_e164=e164,
                country_iso=iso,
                parse_confidence="1.0",
                tags=[],
            )
            db.add(contact)
            db.flush()
            inserted += 1
        else:
            updated += 1
        contact.display_name = (row.get("full_name") or row.get("name") or contact.display_name or "")[:500]
        contact.city = (row.get("city") or contact.city or "")[:120] or None
        contact.preferred_locale = (row.get("preferred_language") or row.get("locale") or contact.preferred_locale or "")[:16] or None
        contact.source = (row.get("contact_source") or row.get("source") or contact.source or "")[:120] or None
        contact.country_iso = iso or contact.country_iso
        tag_raw = row.get("tags") or ""
        if tag_raw:
            contact.tags = [part.strip() for part in tag_raw.split("|") if part.strip()]
        contact.updated_at = utcnow()
        purpose = (row.get("marketing_purpose") or "marketing").strip()[:64] or "marketing"
        kind = classify_consent(row)
        suppressed = active_suppression(db, contact.id) is not None
        if kind == "withdrawn":
            suppress_contact(db, contact, reason="import_withdrawn", actor_id=actor_id, source="import")
            needs_consent += 1
        elif kind == "granted" and not suppressed:
            obtained = _parse_time(row.get("consent_timestamp"))
            db.add(
                CommConsent(
                    contact_id=contact.id,
                    business_unit_id=business_unit_id,
                    purpose=purpose,
                    channel="whatsapp",
                    status="granted",
                    obtained_at=obtained,
                    source=row.get("consent_source"),
                    notice_version=row.get("notice_version") or row.get("consent_source"),
                    proof_uri=row.get("consent_proof_ref"),
                )
            )
            if _purpose_matches(purpose, "marketing"):
                valid_opt_in += 1
            else:
                needs_consent += 1
        else:
            db.add(
                CommConsent(
                    contact_id=contact.id,
                    business_unit_id=business_unit_id,
                    purpose=purpose,
                    channel="whatsapp",
                    status="unconfirmed" if kind == "granted" else "pending",
                    source=row.get("consent_source") or None,
                    proof_uri=row.get("consent_proof_ref") or None,
                    evidence_note="unproven_import" if kind == "granted" else None,
                )
            )
            needs_consent += 1
        if created:
            pass
    job.count_total = len(rows)
    job.inserted = inserted
    job.updated = updated
    job.rejected = rejected
    job.with_valid_marketing_opt_in = valid_opt_in
    job.requires_consent = needs_consent
    job.rejections_json = {"counts": counts, "samples": samples}
    job.status = "finished"
    job.finished_at = utcnow()
    try:
        from app.services.storage import get_storage

        key = f"communications/imports/{job.id}.bin"
        get_storage().upload_bytes(key, data, "application/octet-stream")
        job.file_uri = key
    except Exception:
        job.file_uri = None
    audit(
        db,
        actor_id=actor_id,
        action="contacts.import",
        object_type="import_job",
        object_id=job.id,
        metadata={"total": job.count_total, "inserted": inserted, "rejected": rejected},
    )
    db.flush()
    return job


def exclusion_reason(
    db: Session,
    contact: CommContact,
    *,
    business_unit_id: str,
    purpose: str,
    category: str,
    now: datetime | None = None,
) -> str | None:
    if contact.deleted_at is not None or contact.merged_into_id:
        return "duplicate"
    if not contact.phone_e164 or not contact.country_iso:
        return "invalid_number"
    if contact.blocked or active_suppression(db, contact.id, now):
        return "opted_out"
    if not has_verified_consent(db, contact.id, business_unit_id=business_unit_id, purpose=purpose):
        return "no_consent"
    if category == "marketing" and in_marketing_cooldown(db, contact.id, now):
        return "cooldown"
    return None


def contact_matches(contact: CommContact, filt: dict[str, Any]) -> bool:
    if filt.get("city") and (contact.city or "").lower() != str(filt["city"]).lower():
        return False
    if filt.get("locale") and (contact.preferred_locale or "") != str(filt["locale"]):
        return False
    if filt.get("country_iso") and (contact.country_iso or "") != str(filt["country_iso"]):
        return False
    if filt.get("source") and (contact.source or "") != str(filt["source"]):
        return False
    wanted = filt.get("tags") or []
    if wanted and not set(wanted).issubset(set(contact.tags or [])):
        return False
    return True
